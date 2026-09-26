"""Writes the viewer fixtures in this folder. Not run by the suites; rerun by hand
(python-docx, python-pptx and openpyxl) only when a fixture has to change."""
import datetime
import pathlib
import struct
import zlib

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

here = pathlib.Path(__file__).parent


def png(width, height, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * width for _ in range(height))

    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))

    header = struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


(here / 'picture.png').write_bytes(png(120, 80, (200, 90, 40)))


def pdf():
    # Page 1 uses a Japanese font the file does not embed, so it needs the UniJIS
    # CMap that pdf.js bundles; page 2 uses a standard font, which needs its font data.
    text = '資料ビューアーの試験'.encode('utf-16-be').hex().upper()
    pages = [
        f'BT /F1 28 Tf 72 700 Td <{text}> Tj ET',
        'BT /F2 24 Tf 72 700 Td (Second page in Helvetica) Tj ET',
    ]
    objects = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 5 0 R /Resources << /Font << /F1 7 0 R >> >> >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 6 0 R /Resources << /Font << /F2 9 0 R >> >> >>',
        f'<< /Length {len(pages[0])} >>\nstream\n{pages[0]}\nendstream',
        f'<< /Length {len(pages[1])} >>\nstream\n{pages[1]}\nendstream',
        '<< /Type /Font /Subtype /Type0 /BaseFont /HeiseiMin-W3 /Encoding /UniJIS-UCS2-H /DescendantFonts [8 0 R] >>',
        '<< /Type /Font /Subtype /CIDFontType0 /BaseFont /HeiseiMin-W3 /CIDSystemInfo << /Registry (Adobe) /Ordering (Japan1) /Supplement 2 >> /FontDescriptor 10 0 R >>',
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        '<< /Type /FontDescriptor /FontName /HeiseiMin-W3 /Flags 6 /FontBBox [-123 -257 1001 910] /ItalicAngle 0 /Ascent 723 /Descent -241 /CapHeight 709 /StemV 69 >>',
    ]
    out = b'%PDF-1.4\n'
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f'{number} 0 obj\n{body}\nendobj\n'.encode('latin-1')
    xref = len(out)
    out += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    for offset in offsets:
        out += f'{offset:010d} 00000 n \n'.encode()
    out += f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()
    (here / 'sample.pdf').write_bytes(out)


pdf()

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = '四半期レビュー <script>alert(1)</script>'
slide.placeholders[1].text = 'irori 資料ビューアー'
slide = deck.slides.add_slide(deck.slide_layouts[1])
slide.shapes.title.text = '要点'
frame = slide.placeholders[1].text_frame
frame.text = '売上は 12% 増加'
for line in ['新規顧客 34 社', '解約率 1.8%']:
    paragraph = frame.add_paragraph()
    paragraph.text = line
    paragraph.level = 1
slide.shapes.add_picture(str(here / 'picture.png'), Inches(6), Inches(5), Inches(2))
box = slide.shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(5), Inches(0.6))
run = box.text_frame.paragraphs[0].add_run()
run.text = '赤い太字の注記'
run.font.bold = True
run.font.color.rgb = RGBColor(200, 0, 0)
run.font.size = Pt(20)
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = '表とグラフ'
table = slide.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(4), Inches(1.5)).table
for i, row in enumerate([['地域', 'Q1', 'Q2'], ['東京', '10', '12'], ['大阪', '7', '9']]):
    for j, value in enumerate(row):
        table.cell(i, j).text = value
data = CategoryChartData()
data.categories = ['Q1', 'Q2', 'Q3']
data.add_series('売上', (10, 12, 15))
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4.5), Inches(3), data)
deck.save(here / 'sample.pptx')


def hyperlink(paragraph, url, text):
    part = paragraph.part
    rel = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), rel)
    run = OxmlElement('w:r')
    node = OxmlElement('w:t')
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


doc = Document()
doc.add_heading('議事録 2026-09-26', 0)
doc.add_paragraph('出席者: A, B, C')
doc.add_heading('決定事項', 1)
doc.add_paragraph('PDF と Office を irori で見られるようにする', style='List Bullet')
paragraph = doc.add_paragraph('これは ')
paragraph.add_run('太字').bold = True
paragraph.add_run(' と ')
paragraph.add_run('斜体').italic = True
hyperlink(doc.add_paragraph('外部リンク: '), 'https://example.com/irori', 'example.com')
hyperlink(doc.add_paragraph('危険なリンク: '), 'javascript:alert(1)', 'script link')
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
for (r, c), value in {(0, 0): '項目', (0, 1): '担当', (1, 0): 'ビューアー', (1, 1): 'irori'}.items():
    table.cell(r, c).text = value
doc.add_picture(str(here / 'picture.png'))
doc.add_page_break()
doc.add_paragraph('2ページ目 <img src=x onerror=alert(1)>')
doc.save(here / 'sample.docx')

book = Workbook()
sheet = book.active
sheet.title = '売上'
sheet.append(['日付', '地域', '金額', '比率'])
for i in range(1, 251):
    sheet.append([datetime.date(2026, 9, i % 28 + 1), ['東京', '大阪'][i % 2], i * 1234.5, i / 1000])
for cell in sheet['D'][1:]:
    cell.number_format = '0.0%'
for cell in sheet['C'][1:]:
    cell.number_format = '#,##0'
for cell in sheet['A'][1:]:
    cell.number_format = 'yyyy-mm-dd'
sheet['F1'] = '合計'
sheet['G1'] = 'formula-free total'
sheet.merge_cells('F3:G4')
sheet['F3'] = '結合セル'
book.create_sheet('メモ')['A1'] = '二枚目のシート'
book.save(here / 'sample.xlsx')
